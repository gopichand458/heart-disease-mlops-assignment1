"""Assemble the final Word report (report/MLOps_Assignment1_Report.docx).

Run AFTER training and after you've saved screenshots into screenshots/:
    python report/build_report.py --name "Your Name" --id "2025AE05912" \
        --repo-url https://github.com/<you>/heart-disease-mlops-pipeline

Figures come from reports/figures/ (python -m src.eda), metrics from
models/model_metadata.json. Missing screenshots become visible placeholders.
Edit the text below to reflect your own reasoning - it is a draft, not boilerplate.
"""
import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

FIG = ROOT / "reports" / "figures"
SHOTS = ROOT / "screenshots"
OUT = ROOT / "report" / "MLOps_Assignment1_Report.docx"


def picture(doc, path: Path, caption: str, width=6.0):
    if path.exists():
        doc.add_picture(str(path), width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p = doc.add_paragraph(f"[MISSING IMAGE: {path.name} - add it to {path.parent.name}/]")
        p.runs[0].italic = True
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def metrics_table(doc, meta):
    m = meta.get("metrics", {})
    rows = [("Accuracy", "cv_accuracy_mean", "test_accuracy"),
            ("Precision", "cv_precision_mean", "test_precision"),
            ("Recall", "cv_recall_mean", "test_recall"),
            ("F1", "cv_f1_mean", "test_f1"),
            ("ROC-AUC", "cv_roc_auc_mean", "test_roc_auc")]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ("Metric", "5-fold CV (train)", "Hold-out test")):
        cell.text = text
    for label, cv_key, test_key in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = f"{m[cv_key]:.3f}" if cv_key in m else "n/a"
        cells[2].text = f"{m[test_key]:.3f}" if test_key in m else "n/a"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--name", default="<Your Name>")
    ap.add_argument("--id", default="<Student ID>")
    ap.add_argument("--repo-url", default="https://github.com/<you>/heart-disease-mlops-pipeline")
    args = ap.parse_args()

    meta_path = ROOT / "models" / "model_metadata.json"
    meta = json.loads(meta_path.read_text()) if meta_path.exists() else {}
    best = meta.get("model_name", "<model>")

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("MLOps Assignment 1", 0)
    doc.add_paragraph("End-to-End ML Model Development, CI/CD and Production Deployment\n"
                      "Heart Disease Risk Prediction (UCI Cleveland dataset)")
    doc.add_paragraph(f"Student: {args.name}   |   ID: {args.id}   |   Course: AIMLCZG523")
    doc.add_paragraph(f"Code repository: {args.repo_url}")
    doc.add_page_break()

    doc.add_heading("1. Introduction and problem statement", 1)
    doc.add_paragraph(
        "The goal is to predict whether a patient has heart disease from 13 routine clinical "
        "measurements and to ship that model as a monitored, containerised API. The work "
        "follows an MLOps lifecycle: reproducible data preparation, tracked experiments, "
        "automated testing in CI, a Docker image, a Kubernetes deployment, and runtime "
        "monitoring.")

    doc.add_heading("2. Setup and installation", 1)
    doc.add_paragraph("Tested on the Linux lab VM (Python 3.12, Docker 28, Minikube 1.35).")
    code(doc, "git clone " + args.repo_url + "\ncd heart-disease-mlops-pipeline\n"
              "python -m venv .venv && source .venv/bin/activate\n"
              "pip install -r requirements-dev.txt\n"
              "python scripts/download_data.py     # downloads + cleans the dataset\n"
              "python -m src.eda                   # EDA figures\n"
              "python -m src.train                 # train, track in MLflow, save model\n"
              "pytest                              # unit tests\n"
              "make docker-build && make docker-run\n"
              "bash scripts/deploy_k8s.sh          # minikube deployment + monitoring")

    doc.add_heading("3. Data acquisition and EDA", 1)
    doc.add_paragraph(
        "The dataset is the Cleveland subset of the UCI Heart Disease data (303 patients, 13 "
        "features plus a 0-4 severity label). scripts/download_data.py downloads the file, "
        "adds column names and converts '?' to missing values; src/data.py then coerces "
        "types, binarises the target (0 = no disease, 1-4 = disease) and removes duplicates. "
        "Only 'ca' and 'thal' contain missing values. They are kept as NaN in the cleaned "
        "file and imputed inside the sklearn Pipeline so that imputation statistics are "
        "learned from training folds only, which prevents leakage into cross-validation.")
    picture(doc, FIG / "class_balance.png", "Figure 1 - Class balance", 3.4)
    picture(doc, FIG / "numeric_histograms.png", "Figure 2 - Numeric feature distributions by target")
    picture(doc, FIG / "correlation_heatmap.png", "Figure 3 - Correlation heatmap", 5.0)
    picture(doc, FIG / "categorical_counts.png", "Figure 4 - Categorical features by target")
    doc.add_paragraph("Key observations: [WRITE YOUR OWN: class balance, which features separate "
                      "the classes, outliers in chol/oldpeak, strongest correlations.]")

    doc.add_heading("4. Feature engineering and model development", 1)
    bullets(doc, [
        "Numeric features: median imputation then standardisation.",
        "Categorical features (sex, cp, fbs, restecg, exang, slope, ca, thal): mode "
        "imputation then one-hot encoding (unknown categories ignored at inference).",
        "Engineered feature pct_max_hr = thalach / (220 - age): the heart rate achieved "
        "relative to the age-predicted maximum.",
        "Three models were tuned with GridSearchCV (5-fold stratified, ROC-AUC): Logistic "
        "Regression, Random Forest and Gradient Boosting. The whole pipeline (preprocessing "
        "+ model) is refit inside each fold.",
        "Evaluation: 5-fold CV on an 80% training split plus an untouched 20% hold-out test set; "
        "metrics are accuracy, precision, recall, F1 and ROC-AUC.",
        f"Selection rule: highest cross-validated ROC-AUC. Selected model: {best}.",
    ])
    metrics_table(doc, meta)
    doc.add_paragraph("")
    doc.add_paragraph("Model selection discussion: [WRITE YOUR OWN: why the winner won, "
                      "recall vs precision trade-off for a screening use case.]")

    doc.add_heading("5. Experiment tracking (MLflow)", 1)
    doc.add_paragraph(
        "Each training invocation creates a parent run 'model-comparison' with one nested run "
        "per model. Logged: dataset size and class rate, best hyper-parameters, CV and test "
        "metrics, ROC curve, confusion matrix, feature-importance plot, grid-search results "
        "CSV and the fitted model with signature and input example. The parent run stores the "
        "model comparison chart and the final deployable artifacts.")
    picture(doc, SHOTS / "03_mlflow_runs.png", "Figure 5 - MLflow experiment runs")
    picture(doc, SHOTS / "04_mlflow_compare.png", "Figure 6 - MLflow run comparison / metrics")

    doc.add_heading("6. Model packaging and reproducibility", 1)
    bullets(doc, [
        "A single sklearn Pipeline (preprocessing + model) is saved as models/best_model.joblib "
        "and in MLflow format (models/mlflow_model) so raw patient fields can be scored directly.",
        "models/model_metadata.json records the version, git SHA, library versions, features "
        "and metrics.",
        "requirements.txt (training), requirements-dev.txt (testing/docs), requirements-serve.txt "
        "(minimal serving image). Random seeds are fixed.",
    ])

    doc.add_heading("7. CI/CD pipeline and automated testing", 1)
    doc.add_paragraph(
        "tests/ contains unit tests for data cleaning, the feature pipeline, modelling and "
        "metrics, and the API (validation, prediction schema, metrics endpoint, request "
        "logging). The GitHub Actions workflow (.github/workflows/ci.yml) runs flake8, pytest "
        "with coverage, trains the models, and uploads test reports, the model and MLflow runs "
        "as artifacts; a second job builds the Docker image and smoke-tests /predict. Any lint "
        "or test failure stops the pipeline.")
    picture(doc, SHOTS / "10_pytest_pass.png", "Figure 7 - Local pytest run")
    picture(doc, SHOTS / "01_ci_run.png", "Figure 8 - Successful GitHub Actions run")
    picture(doc, SHOTS / "02_ci_artifacts.png", "Figure 9 - Workflow artifacts")
    picture(doc, SHOTS / "11_ci_failure_demo.png", "Figure 10 - Pipeline failing on a deliberate error")

    doc.add_heading("8. Containerisation", 1)
    doc.add_paragraph(
        "The API (FastAPI) exposes POST /predict (JSON in; prediction, confidence, probability "
        "and risk band out), GET /health and GET /metrics. Inputs are validated with Pydantic "
        "ranges (invalid data returns HTTP 422). The image uses python:3.12-slim, runs as a "
        "non-root user and has a health check.")
    code(doc, "docker build --build-arg SKLEARN_VERSION=$(python -c 'import sklearn;"
              "print(sklearn.__version__)') -t heart-disease-api:1.0 .\n"
              "docker run -d -p 8000:8000 heart-disease-api:1.0\n"
              "bash scripts/smoke_test.sh")
    picture(doc, SHOTS / "05_docker_run_predict.png", "Figure 11 - Container build/run and sample prediction")

    doc.add_heading("9. Production deployment (Kubernetes)", 1)
    doc.add_paragraph(
        "The image is deployed to Minikube with k8s/deployment.yaml (2 replicas, readiness and "
        "liveness probes, resource requests/limits) and k8s/service.yaml (type LoadBalancer, "
        "exposed with 'minikube tunnel'); k8s/ingress.yaml is an optional ingress alternative.")
    picture(doc, SHOTS / "06_k8s_pods_svc.png", "Figure 12 - Pods and service")
    picture(doc, SHOTS / "07_api_swagger_k8s.png", "Figure 13 - Endpoint verified through the service")

    doc.add_heading("10. Monitoring and logging", 1)
    doc.add_paragraph(
        "Every request is logged as one JSON line (id, method, path, status, latency) to stdout "
        "and logs/api.log. Prometheus scrapes /metrics (api_requests_total, "
        "api_request_latency_seconds, predictions_total by risk band, prediction_errors_total) "
        "and Grafana shows a provisioned dashboard.")
    picture(doc, SHOTS / "08_prometheus_targets.png", "Figure 14 - Prometheus target UP")
    picture(doc, SHOTS / "09_grafana_dashboard.png", "Figure 15 - Grafana dashboard")

    doc.add_heading("11. Architecture", 1)
    picture(doc, ROOT / "docs" / "architecture.png", "Figure 16 - End-to-end architecture", 6.5)

    doc.add_heading("12. Conclusion, limitations and future work", 1)
    bullets(doc, [
        "The small dataset (303 rows) limits the stability of the metrics; confidence intervals "
        "would be wide.",
        "Future work: model registry and promotion stages, data/prediction drift monitoring, "
        "canary deployments, Helm packaging, HPA autoscaling.",
        "This model is an educational exercise and not a medical device.",
    ])
    doc.add_heading("References", 2)
    bullets(doc, ["UCI Machine Learning Repository - Heart Disease dataset.",
                  "MLflow, scikit-learn, FastAPI, Prometheus and Grafana documentation.",
                  f"Repository: {args.repo_url}"])

    doc.save(OUT)
    print("Saved", OUT)


if __name__ == "__main__":
    main()
